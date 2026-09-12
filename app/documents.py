"""Shared PDF and Excel generators for Aurora documents."""
from __future__ import annotations

from datetime import datetime
from pathlib import Path
from html import escape
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from .database import Database


class DocumentEngine:
    def __init__(self, db: Database, root: Path): self.db, self.root = db, root

    def _pdf_styles(self):
        styles=getSampleStyleSheet()
        return {
            'title': ParagraphStyle('AuroraTitle', parent=styles['Title'], fontName='Helvetica-Bold', fontSize=21, leading=24, textColor=colors.HexColor('#17171A'), spaceAfter=2),
            'eyebrow': ParagraphStyle('AuroraEyebrow', parent=styles['BodyText'], fontName='Helvetica-Bold', fontSize=8.5, leading=10, textColor=colors.HexColor('#5B5FEF')),
            'body': ParagraphStyle('AuroraBody', parent=styles['BodyText'], fontName='Helvetica', fontSize=9.2, leading=13, textColor=colors.HexColor('#17171A')),
            'small': ParagraphStyle('AuroraSmall', parent=styles['BodyText'], fontName='Helvetica', fontSize=8, leading=10.5, textColor=colors.HexColor('#6F7078')),
            'cell': ParagraphStyle('AuroraCell', parent=styles['BodyText'], fontName='Helvetica', fontSize=7.9, leading=10, textColor=colors.HexColor('#17171A')),
            'cell_right': ParagraphStyle('AuroraCellRight', parent=styles['BodyText'], fontName='Helvetica', fontSize=7.9, leading=10, textColor=colors.HexColor('#17171A'), alignment=TA_RIGHT),
            'total_label': ParagraphStyle('AuroraTotalLabel', parent=styles['BodyText'], fontName='Helvetica', fontSize=8.5, leading=11, textColor=colors.HexColor('#6F7078'), alignment=TA_RIGHT),
            'total_value': ParagraphStyle('AuroraTotalValue', parent=styles['BodyText'], fontName='Helvetica-Bold', fontSize=10, leading=12, textColor=colors.HexColor('#17171A'), alignment=TA_RIGHT),
            'grand_value': ParagraphStyle('AuroraGrandValue', parent=styles['BodyText'], fontName='Helvetica-Bold', fontSize=15, leading=17, textColor=colors.HexColor('#5B5FEF'), alignment=TA_RIGHT),
        }

    @staticmethod
    def _esc(value):
        return escape(str(value or '')).replace('\n','<br/>')

    @staticmethod
    def _money(value):
        return f"INR {float(value or 0):,.2f}"

    @staticmethod
    def _footer(canvas, doc):
        canvas.saveState(); width,_=A4; canvas.setStrokeColor(colors.HexColor('#E6E6EA')); canvas.setLineWidth(.5); canvas.line(15*mm,12*mm,width-15*mm,12*mm); canvas.setFillColor(colors.HexColor('#6F7078')); canvas.setFont('Helvetica',7.5); canvas.drawString(15*mm,7.5*mm,'AURORA CLOUD · NEXUS · by hash'); canvas.drawRightString(width-15*mm,7.5*mm,f'Page {doc.page}'); canvas.restoreState()

    def _header(self, styles, doc):
        mark=Table([[Paragraph('A', ParagraphStyle('Mark',fontName='Helvetica-Bold',fontSize=15,leading=18,textColor=colors.white,alignment=TA_CENTER))]],colWidths=[14*mm],rowHeights=[14*mm])
        mark.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,-1),colors.HexColor('#5B5FEF')),('VALIGN',(0,0),(-1,-1),'MIDDLE'),('LEFTPADDING',(0,0),(-1,-1),0),('RIGHTPADDING',(0,0),(-1,-1),0),('TOPPADDING',(0,0),(-1,-1),0),('BOTTOMPADDING',(0,0),(-1,-1),0)]))
        brand=Table([[mark, [Paragraph('AURORA CLOUD',styles['title']),Paragraph('NEXUS · BUSINESS MANAGEMENT',styles['small'])]]],colWidths=[17*mm,93*mm])
        brand.setStyle(TableStyle([('VALIGN',(0,0),(-1,-1),'MIDDLE'),('LEFTPADDING',(0,0),(-1,-1),0),('RIGHTPADDING',(0,0),(-1,-1),0),('TOPPADDING',(0,0),(-1,-1),0),('BOTTOMPADDING',(0,0),(-1,-1),0)]))
        meta=Table([[Paragraph(doc['kind'].replace('_',' ').upper(),styles['eyebrow'])],[Paragraph(self._esc(doc['number']),styles['title'])],[Paragraph(f"Issue date: {self._esc(doc['issue_date'])}",styles['small'])],[Paragraph(f"Status: <b>{self._esc(doc['status'])}</b>",styles['small'])]],colWidths=[70*mm])
        meta.setStyle(TableStyle([('ALIGN',(0,0),(-1,-1),'RIGHT'),('LEFTPADDING',(0,0),(-1,-1),0),('RIGHTPADDING',(0,0),(-1,-1),0),('TOPPADDING',(0,0),(-1,-1),0),('BOTTOMPADDING',(0,0),(-1,-1),1)]))
        h=Table([[brand,meta]],colWidths=[110*mm,70*mm]); h.setStyle(TableStyle([('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),0),('RIGHTPADDING',(0,0),(-1,-1),0),('TOPPADDING',(0,0),(-1,-1),0),('BOTTOMPADDING',(0,0),(-1,-1),0)])); return h

    def _party(self, styles, doc):
        left=[Paragraph('BILL TO',styles['eyebrow']),Paragraph(self._esc(doc['company_name'] or '—'),ParagraphStyle('Party',parent=styles['body'],fontName='Helvetica-Bold',fontSize=11,leading=14)),Paragraph(self._esc(doc.get('address') or ''),styles['small']),Paragraph(self._esc(doc.get('email') or ''),styles['small'])]
        if doc.get('tax_id'): left.append(Paragraph(f"GST / Tax ID: {self._esc(doc['tax_id'])}",styles['small']))
        right=[Paragraph('DOCUMENT',styles['eyebrow']),Paragraph(f"Number: <b>{self._esc(doc['number'])}</b>",styles['small']),Paragraph(f"Date: {self._esc(doc['issue_date'])}",styles['small']),Paragraph(f"Revision: {self._esc(doc.get('revision',1))}",styles['small'])]
        t=Table([[left,right]],colWidths=[112*mm,68*mm]); t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,-1),colors.HexColor('#F7F7FA')),('BOX',(0,0),(-1,-1),.6,colors.HexColor('#E6E6EA')),('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),9),('RIGHTPADDING',(0,0),(-1,-1),9),('TOPPADDING',(0,0),(-1,-1),9),('BOTTOMPADDING',(0,0),(-1,-1),9)])); return t

    def _lines(self, styles, lines):
        data=[[Paragraph(x,ParagraphStyle('th'+str(i),fontName='Helvetica-Bold',fontSize=7.7,leading=9,textColor=colors.white)) for i,x in enumerate(['DESCRIPTION','QTY','UNIT','RATE','DISC.','TAX','AMOUNT'])]]
        for l in lines:
            base=float(l['quantity'] or 0)*float(l['rate'] or 0)-float(l['discount'] or 0); amount=base*(1+float(l['tax_rate'] or 0)/100)
            data.append([Paragraph(self._esc(l['description']),styles['cell']),Paragraph(f"{float(l['quantity']):g}",styles['cell_right']),Paragraph(self._esc(l['unit']),styles['cell']),Paragraph(self._money(l['rate']),styles['cell_right']),Paragraph(self._money(l['discount']),styles['cell_right']),Paragraph(f"{float(l['tax_rate'] or 0):g}%",styles['cell_right']),Paragraph(self._money(amount),styles['cell_right'])])
        t=Table(data,colWidths=[62*mm,13*mm,18*mm,25*mm,20*mm,14*mm,28*mm],repeatRows=1); t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#5B5FEF')),('BOX',(0,0),(-1,-1),.6,colors.HexColor('#E6E6EA')),('INNERGRID',(0,1),(-1,-1),.35,colors.HexColor('#E6E6EA')),('VALIGN',(0,0),(-1,-1),'MIDDLE'),('LEFTPADDING',(0,0),(-1,-1),6),('RIGHTPADDING',(0,0),(-1,-1),6),('TOPPADDING',(0,0),(-1,-1),6),('BOTTOMPADDING',(0,0),(-1,-1),6)]))
        for r in range(2,len(data)+1):
            if r%2==0: t.setStyle(TableStyle([('BACKGROUND',(0,r-1),(-1,r-1),colors.HexColor('#FBFBFC'))]))
        return t

    def _totals(self, styles, doc):
        rows=[[Paragraph('Subtotal',styles['total_label']),Paragraph(self._money(doc['subtotal']),styles['total_value'])],[Paragraph('Discount',styles['total_label']),Paragraph(self._money(doc.get('discount',0)),styles['total_value'])],[Paragraph('Tax',styles['total_label']),Paragraph(self._money(doc['tax']),styles['total_value'])],[Paragraph('TOTAL',ParagraphStyle('GrandLabel',parent=styles['total_label'],fontName='Helvetica-Bold',textColor=colors.HexColor('#17171A'))),Paragraph(self._money(doc['total']),styles['grand_value'])]]
        t=Table(rows,colWidths=[40*mm,45*mm],hAlign='RIGHT'); t.setStyle(TableStyle([('LINEABOVE',(0,3),(-1,3),1.2,colors.HexColor('#5B5FEF')),('BACKGROUND',(0,3),(-1,3),colors.HexColor('#EEF0FF')),('LEFTPADDING',(0,0),(-1,-1),6),('RIGHTPADDING',(0,0),(-1,-1),6),('TOPPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),5)])); return Table([[Spacer(1,1),t]],colWidths=[95*mm,85*mm],style=TableStyle([('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),0),('RIGHTPADDING',(0,0),(-1,-1),0),('TOPPADDING',(0,0),(-1,-1),0),('BOTTOMPADDING',(0,0),(-1,-1),0)]))

    def pdf(self, doc_id: int) -> Path:
        rows=self.db.rows('SELECT d.*,c.company_name,c.address,c.email,c.tax_id,c.phone FROM documents d LEFT JOIN customers c ON c.id=d.customer_id WHERE d.id=?',(doc_id,))
        if not rows: raise ValueError('Document not found.')
        doc=dict(rows[0]); lines=self.db.rows('SELECT * FROM document_lines WHERE document_id=? ORDER BY id',(doc_id,)); folder={'QUOTE':'Quotes','INVOICE':'Invoices','PURCHASE_ORDER':'Purchase Orders'}[doc['kind']]; path=self.root/'Documents'/folder/f"{doc['number']}.pdf"; path.parent.mkdir(parents=True,exist_ok=True); st=self._pdf_styles()
        story=[self._header(st,doc),Spacer(1,6*mm),self._party(st,doc),Spacer(1,7*mm),Paragraph('LINE ITEMS',st['eyebrow']),Spacer(1,2*mm),self._lines(st,lines),Spacer(1,5*mm),self._totals(st,doc)]
        blocks=[]
        if doc.get('notes'): blocks.append([Paragraph('NOTES',st['eyebrow']),Paragraph(self._esc(doc['notes']),st['small'])])
        if doc.get('terms'): blocks.append([Paragraph('TERMS',st['eyebrow']),Paragraph(self._esc(doc['terms']),st['small'])])
        if blocks:
            nt=Table(blocks,colWidths=[30*mm,150*mm]); nt.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,-1),colors.HexColor('#F7F7FA')),('BOX',(0,0),(-1,-1),.6,colors.HexColor('#E6E6EA')),('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),8),('RIGHTPADDING',(0,0),(-1,-1),8),('TOPPADDING',(0,0),(-1,-1),7),('BOTTOMPADDING',(0,0),(-1,-1),7)])); story += [Spacer(1,6*mm),nt]
        story += [Spacer(1,8*mm),Paragraph('Thank you for your business.',ParagraphStyle('Thanks',parent=st['body'],fontName='Helvetica-Bold',fontSize=9.5)),Paragraph('Generated by AURORA CLOUD — NEXUS',st['small'])]
        SimpleDocTemplate(str(path),pagesize=A4,rightMargin=15*mm,leftMargin=15*mm,topMargin=14*mm,bottomMargin=16*mm,title=f"{doc['kind']} {doc['number']}",author='AURORA CLOUD — NEXUS').build(story,onFirstPage=self._footer,onLaterPages=self._footer)
        with self.db.transaction() as con: self.db.log(con,'PDF generated','document',doc_id,str(path))
        return path

    def receipt_pdf(self, payment_id: int) -> Path:
        rows=self.db.rows('SELECT p.*,d.number,d.total invoice_total,c.company_name,c.email,c.address FROM payments p JOIN documents d ON d.id=p.invoice_id LEFT JOIN customers c ON c.id=d.customer_id WHERE p.id=?',(payment_id,))
        if not rows: raise ValueError('Payment not found.')
        pay=dict(rows[0]); path=self.root/'Documents'/'Receipts'/f"REC-{payment_id:05d}.pdf"; path.parent.mkdir(parents=True,exist_ok=True); st=self._pdf_styles()
        mark=Paragraph('A',ParagraphStyle('ReceiptMark',fontName='Helvetica-Bold',fontSize=15,textColor=colors.white,alignment=TA_CENTER)); header=Table([[mark,Paragraph('AURORA CLOUD<br/><font size="8">NEXUS · PAYMENT RECEIPT</font>',st['title']),Paragraph('RECEIPT',st['eyebrow'])]],colWidths=[14*mm,111*mm,55*mm]); header.setStyle(TableStyle([('BACKGROUND',(0,0),(0,0),colors.HexColor('#5B5FEF')),('ALIGN',(2,0),(2,0),'RIGHT'),('VALIGN',(0,0),(-1,-1),'MIDDLE'),('LEFTPADDING',(0,0),(-1,-1),0),('RIGHTPADDING',(0,0),(-1,-1),0)]))
        info=Table([[Paragraph('RECEIPT NUMBER',st['eyebrow']),Paragraph('RECEIVED FROM',st['eyebrow'])],[Paragraph(f"REC-{payment_id:05d}",st['title']),Paragraph(self._esc(pay['company_name'] or '—'),ParagraphStyle('RP',parent=st['body'],fontName='Helvetica-Bold',fontSize=12))],[Paragraph(f"Payment date: {self._esc(pay['payment_date'])}",st['small']),Paragraph(self._esc(pay['email'] or ''),st['small'])],[Paragraph(f"Invoice: {self._esc(pay['number'])}",st['small']),Paragraph(self._esc(pay['address'] or ''),st['small'])]],colWidths=[75*mm,105*mm]); info.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,-1),colors.HexColor('#F7F7FA')),('BOX',(0,0),(-1,-1),.7,colors.HexColor('#E6E6EA')),('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),9),('RIGHTPADDING',(0,0),(-1,-1),9),('TOPPADDING',(0,0),(-1,-1),8),('BOTTOMPADDING',(0,0),(-1,-1),8)]))
        box=Table([[Paragraph('AMOUNT RECEIVED',st['eyebrow']),Paragraph(self._money(pay['amount']),st['grand_value'])],[Paragraph('METHOD',st['small']),Paragraph(self._esc(pay['method']),st['body'])],[Paragraph('REFERENCE',st['small']),Paragraph(self._esc(pay['reference'] or '—'),st['body'])],[Paragraph('INVOICE TOTAL',st['small']),Paragraph(self._money(pay['invoice_total']),st['body'])]],colWidths=[80*mm,100*mm]); box.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#EEF0FF')),('BOX',(0,0),(-1,-1),.7,colors.HexColor('#E6E6EA')),('VALIGN',(0,0),(-1,-1),'MIDDLE'),('ALIGN',(1,0),(1,0),'RIGHT'),('LEFTPADDING',(0,0),(-1,-1),9),('RIGHTPADDING',(0,0),(-1,-1),9),('TOPPADDING',(0,0),(-1,-1),8),('BOTTOMPADDING',(0,0),(-1,-1),8)]))
        story=[header,Spacer(1,8*mm),info,Spacer(1,8*mm),box,Spacer(1,8*mm),Paragraph('Thank you — payment received.',ParagraphStyle('Thanks2',parent=st['body'],fontName='Helvetica-Bold',fontSize=11))]
        if pay.get('notes'): story += [Spacer(1,5*mm),Paragraph('NOTES',st['eyebrow']),Paragraph(self._esc(pay['notes']),st['small'])]
        SimpleDocTemplate(str(path),pagesize=A4,rightMargin=15*mm,leftMargin=15*mm,topMargin=14*mm,bottomMargin=16*mm,title=f"Receipt REC-{payment_id:05d}",author='AURORA CLOUD — NEXUS').build(story,onFirstPage=self._footer,onLaterPages=self._footer)
        with self.db.transaction() as con: self.db.log(con,'Receipt PDF generated','payment',payment_id,str(path))
        return path

    def excel(self, title: str, headers: list[str], rows: list[tuple], target: Path) -> Path:
        wb = Workbook()
        ws = wb.active
        ws.title = title[:31]
        ws.append(headers)
        for r in rows:
            ws.append(list(r))
        fill = PatternFill("solid", fgColor="4F46E5")
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = fill
            c.alignment = Alignment(horizontal="center")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(max(len(str(x.value or "")) for x in col) + 2, 40)
        target.parent.mkdir(exist_ok=True)
        wb.save(target)
        return target
